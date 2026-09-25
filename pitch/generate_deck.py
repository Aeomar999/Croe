#!/usr/bin/env python3
"""
Generate Croe Investor & Stakeholder Pitch Deck (.pptx)
Adheres to Croe Design System ("Calm over confrontation", 16:9 widescreen, trademark 'croe.' with green dot)
Minimalist, whitespace-heavy aesthetic.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Paths
ASSETS_DIR = r"c:\Users\Jerry\Desktop\PROJECT 2026\Croe\pitch\assets"
OUTPUT_PPTX = r"c:\Users\Jerry\Desktop\PROJECT 2026\Croe\pitch\Croe_Pitch_Deck.pptx"

# Colors
C_DARK_BG = RGBColor(15, 23, 42)       # #0F172A Slate 900
C_LIGHT_BG = RGBColor(255, 255, 255)   # #FFFFFF Pure White
C_OFF_WHITE = RGBColor(248, 250, 252)  # #F8FAFC Slate 50
C_TEXT_HEAD = RGBColor(15, 23, 42)     # #0F172A
C_TEXT_BODY = RGBColor(51, 65, 85)     # #334155
C_TEXT_MUTED = RGBColor(100, 116, 139) # #64748B
C_BLUE = RGBColor(37, 99, 235)         # #2563EB
C_GREEN = RGBColor(31, 193, 107)       # #1FC16B Croe Green
C_LINE = RGBColor(226, 232, 240)       # #E2E8F0
C_WHITE = RGBColor(255, 255, 255)

FONT_TITLE = "Segoe UI"
FONT_BODY = "Segoe UI"

def set_slide_background(slide, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_logo(slide, is_dark=False):
    logo_path = os.path.join(ASSETS_DIR, "croe_wordmark_light.png" if is_dark else "croe_wordmark_dark.png")
    if os.path.exists(logo_path):
        slide.shapes.add_picture(logo_path, Inches(11.2), Inches(0.6), width=Inches(1.2))

def add_header(slide, kicker, headline, is_dark=False):
    # Kicker
    tx_kicker = slide.shapes.add_textbox(Inches(1.2), Inches(0.8), Inches(9.0), Inches(0.4))
    tf_k = tx_kicker.text_frame
    tf_k.word_wrap = True
    p_k = tf_k.paragraphs[0]
    p_k.text = kicker.upper()
    p_k.font.name = FONT_BODY
    p_k.font.size = Pt(12)
    p_k.font.bold = True
    p_k.font.color.rgb = C_GREEN if is_dark else C_BLUE

    # Headline
    tx_head = slide.shapes.add_textbox(Inches(1.15), Inches(1.1), Inches(9.5), Inches(0.8))
    tf_h = tx_head.text_frame
    tf_h.word_wrap = True
    p_h = tf_h.paragraphs[0]
    p_h.text = headline
    p_h.font.name = FONT_TITLE
    p_h.font.size = Pt(32)
    p_h.font.bold = True
    p_h.font.color.rgb = C_WHITE if is_dark else C_TEXT_HEAD

    # Subtle divider line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(2.2), Inches(10.933), Pt(1))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(51, 65, 85) if is_dark else C_LINE
    line.line.fill.background()

    add_logo(slide, is_dark)

def add_notes(slide, notes_text):
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = notes_text

# Initialize Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# ==============================================================================
# SLIDE 1: Title Slide (Dark Theme)
# ==============================================================================
s1 = prs.slides.add_slide(blank_layout)
set_slide_background(s1, C_DARK_BG)

logo_path = os.path.join(ASSETS_DIR, "croe_wordmark_light.png")
if os.path.exists(logo_path):
    s1.shapes.add_picture(logo_path, Inches(1.2), Inches(2.2), width=Inches(3.5))

s1_box = s1.shapes.add_textbox(Inches(1.15), Inches(3.2), Inches(10.5), Inches(2.2))
tf1 = s1_box.text_frame
tf1.word_wrap = True

p1 = tf1.paragraphs[0]
p1.text = "Escrow for social commerce, on Mobile Money."
p1.font.name = FONT_TITLE
p1.font.size = Pt(44)
p1.font.bold = True
p1.font.color.rgb = C_WHITE

p2 = tf1.add_paragraph()
p2.space_before = Pt(20)
p2.text = "Eliminating the 'you send first' standoff that kills WhatsApp & Instagram sales in Ghana."
p2.font.name = FONT_BODY
p2.font.size = Pt(22)
p2.font.color.rgb = RGBColor(148, 163, 184)

meta_tx = s1.shapes.add_textbox(Inches(1.15), Inches(6.0), Inches(10.2), Inches(0.5))
tf_meta = meta_tx.text_frame
p_meta = tf_meta.paragraphs[0]
p_meta.text = "Founding Team  ·  founders@croe.app  ·  Accra, Ghana  ·  Pre-Seed / Pilot Stage"
p_meta.font.name = FONT_BODY
p_meta.font.size = Pt(14)
p_meta.font.color.rgb = C_GREEN

add_notes(s1, "One sentence, then move: 'Croe makes it safe to buy and sell on WhatsApp and Instagram in Ghana, by holding the buyer's Mobile Money until the goods actually arrive.'")

# ==============================================================================
# SLIDE 2: The Problem
# ==============================================================================
s2 = prs.slides.add_slide(blank_layout)
set_slide_background(s2, C_LIGHT_BG)
add_header(s2, "01 / The Friction Point", "Every social commerce sale starts with a standoff.")

# Column 1: The Chat
tx_chat = s2.shapes.add_textbox(Inches(1.2), Inches(2.8), Inches(4.5), Inches(3.5))
tf_chat = tx_chat.text_frame
tf_chat.word_wrap = True

pc_head = tf_chat.paragraphs[0]
pc_head.text = "The Deadlock"
pc_head.font.bold = True
pc_head.font.size = Pt(18)
pc_head.font.color.rgb = C_TEXT_HEAD

pc_v = tf_chat.add_paragraph()
pc_v.space_before = Pt(20)
pc_v.text = "Vendor:\n\"Please send the MoMo first so I can dispatch the rider.\""
pc_v.font.size = Pt(16)
pc_v.font.color.rgb = C_TEXT_BODY

pc_b = tf_chat.add_paragraph()
pc_b.space_before = Pt(20)
pc_b.text = "Buyer:\n\"No - send the goods first, then I'll send the money.\""
pc_b.font.size = Pt(16)
pc_b.font.color.rgb = C_TEXT_BODY
pc_b.font.bold = False

pc_res = tf_chat.add_paragraph()
pc_res.space_before = Pt(30)
pc_res.text = "Result: 25%–30% of sales abandoned right here."
pc_res.font.bold = True
pc_res.font.size = Pt(16)
pc_res.font.color.rgb = C_GREEN

# Divider line between columns
line = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.4), Inches(2.8), Pt(1), Inches(3.5))
line.fill.solid()
line.fill.fore_color.rgb = C_LINE
line.line.fill.background()

# Column 2: Buyer Trap
tx_b = s2.shapes.add_textbox(Inches(7.2), Inches(2.8), Inches(5.0), Inches(1.5))
tf_b = tx_b.text_frame
tf_b.word_wrap = True
pb_h = tf_b.paragraphs[0]
pb_h.text = "The Buyer's Trap"
pb_h.font.bold = True
pb_h.font.size = Pt(18)
pb_h.font.color.rgb = RGBColor(220, 38, 38)
pb_b = tf_b.add_paragraph()
pb_b.space_before = Pt(12)
pb_b.text = "Buyers get scammed with zero recourse. No enforceable receipt, no neutral arbiter, and no way to reverse a Mobile Money transfer once initiated."
pb_b.font.size = Pt(15)
pb_b.font.color.rgb = C_TEXT_BODY

# Column 3: Vendor Dilemma
tx_v = s2.shapes.add_textbox(Inches(7.2), Inches(4.8), Inches(5.0), Inches(1.5))
tf_v = tx_v.text_frame
tf_v.word_wrap = True
pv_h = tf_v.paragraphs[0]
pv_h.text = "The Vendor's Dilemma"
pv_h.font.bold = True
pv_h.font.size = Pt(18)
pv_h.font.color.rgb = C_BLUE
pv_b = tf_v.add_paragraph()
pv_b.space_before = Pt(12)
pv_b.text = "Vendors eat cash-on-delivery rejections, pay wasted dispatch rider fees, and battle fake payment screenshots. Trusting a stranger means risking stock."
pv_b.font.size = Pt(15)
pv_b.font.color.rgb = C_TEXT_BODY

add_notes(s2, "Make this personal: everyone who has bought or sold on Instagram in Ghana has experienced this standoff. This is a STRUCTURAL failure of social commerce, not bad luck.")

# ==============================================================================
# SLIDE 3: Why This Persists
# ==============================================================================
s3 = prs.slides.add_slide(blank_layout)
set_slide_background(s3, C_LIGHT_BG)
add_header(s3, "02 / Market Reality", "The payment rail works. The trust layer does not exist.")

cards_s3 = [
    ("MoMo is Universal", "MTN, Telecel, AirtelTigo", "20M+ registered Mobile Money accounts in Ghana. Almost every merchant and buyer already carries a digital wallet.", C_BLUE),
    ("MoMo is Final by Design", "Irreversible P2P Transfers", "Once sent, MoMo is gone. No chargebacks, no disputes. What makes it great for sending money to family makes it catastrophic for trade between strangers.", RGBColor(220, 38, 38)),
    ("Commerce Absorbed the Cost", "Fraud as an Overhead", "Card networks solved stranger commerce with chargebacks 40 years ago. Mobile money never got that layer. Social commerce absorbed fraud as a cost of doing business.", C_GREEN)
]

for idx, (title, sub, body, accent_col) in enumerate(cards_s3):
    left = Inches(1.2 + idx * 3.7)
    
    # Subtitle pill-like text
    tx_sub = s3.shapes.add_textbox(left, Inches(2.8), Inches(3.2), Inches(0.4))
    p = tx_sub.text_frame.paragraphs[0]
    p.text = sub.upper()
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = accent_col
    
    tx_body = s3.shapes.add_textbox(left, Inches(3.2), Inches(3.3), Inches(3.0))
    tf = tx_body.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.bold = True
    p1.font.size = Pt(20)
    p1.font.color.rgb = C_TEXT_HEAD
    p2 = tf.add_paragraph()
    p2.space_before = Pt(16)
    p2.text = body
    p2.font.size = Pt(15)
    p2.font.color.rgb = C_TEXT_BODY

add_notes(s3, "MoMo's finality is a FEATURE for peer transfers, and a CATASTROPHE for commerce between strangers. That missing trust layer is Croe.")

# ==============================================================================
# SLIDE 4: The Solution
# ==============================================================================
s4 = prs.slides.add_slide(blank_layout)
set_slide_background(s4, C_LIGHT_BG)
add_header(s4, "03 / The Croe Solution", "Croe holds the money until the goods arrive.")

steps = [
    ("01", "LINK", "Vendor generates link", "Shared directly in the WhatsApp / Instagram chat where the sale is happening."),
    ("02", "PAY", "Buyer pays via MoMo", "Standard USSD prompt. Funds are secured in escrow - vendor sees it, but cannot touch it."),
    ("03", "SHIP", "Vendor dispatches", "Vendor dispatches goods against guaranteed money instead of an empty promise."),
    ("04", "RELEASE", "Delivery confirmed", "Buyer confirms delivery in app; vendor is paid automatically, less 2.5% commission.")
]

for idx, (num, action, title, desc) in enumerate(steps):
    left = Inches(1.2 + idx * 2.7)
    
    tx_num = s4.shapes.add_textbox(left, Inches(2.8), Inches(2.4), Inches(0.5))
    pn = tx_num.text_frame.paragraphs[0]
    pn.text = f"{num} - {action}"
    pn.font.bold = True
    pn.font.size = Pt(14)
    pn.font.color.rgb = C_GREEN
    
    tx = s4.shapes.add_textbox(left, Inches(3.4), Inches(2.4), Inches(3.0))
    tf = tx.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.bold = True
    p1.font.size = Pt(18)
    p1.font.color.rgb = C_TEXT_HEAD
    p2 = tf.add_paragraph()
    p2.space_before = Pt(12)
    p2.text = desc
    p2.font.size = Pt(14)
    p2.font.color.rgb = C_TEXT_BODY

add_notes(s4, "Emphasize ZERO behavior change for the buyer. No new app needed to pay; standard MoMo USSD prompt. The vendor sees the money is locked but cannot touch it.")

# ==============================================================================
# SLIDE 5: Product
# ==============================================================================
s5 = prs.slides.add_slide(blank_layout)
set_slide_background(s5, C_OFF_WHITE)
add_header(s5, "04 / Product Architecture", "It's built. Not a prototype.")

tx_l5 = s5.shapes.add_textbox(Inches(1.2), Inches(2.8), Inches(5.5), Inches(4.2))
tf_l5 = tx_l5.text_frame
tf_l5.word_wrap = True

p = tf_l5.paragraphs[0]
p.text = "Engineered for Financial Rigor"
p.font.bold = True
p.font.size = Pt(20)
p.font.color.rgb = C_BLUE

items_l5 = [
    "14-state escrow machine with append-only financial ledger",
    "UPDATE / DELETE permanently revoked at PostgreSQL role level",
    "Raw-body HMAC verification with 300s replay window defense",
    "Distributed Redis fast dedup + durable inbox queue",
    "Tiered KYC verification (Tier 0 phone OTP to Tier 2 EDD)",
    "SHA-256 cryptographic media hashing catching recycled fraud",
    "Daily automated 3-way reconciliation (bank / ledger / provider)"
]

for item in items_l5:
    p = tf_l5.add_paragraph()
    p.space_before = Pt(12)
    p.text = f"•  {item}"
    p.font.size = Pt(14)
    p.font.color.rgb = C_TEXT_BODY

# Right side stats
tx_stat = s5.shapes.add_textbox(Inches(7.6), Inches(2.8), Inches(4.5), Inches(1.5))
tf_stat = tx_stat.text_frame
p_stat = tf_stat.paragraphs[0]
p_stat.text = "354"
p_stat.font.name = FONT_TITLE
p_stat.font.size = Pt(72)
p_stat.font.bold = True
p_stat.font.color.rgb = C_GREEN

p_stat_sub = tf_stat.add_paragraph()
p_stat_sub.text = "Automated tests passing in CI/CD"
p_stat_sub.font.bold = True
p_stat_sub.font.size = Pt(16)
p_stat_sub.font.color.rgb = C_TEXT_HEAD

tx_r5_body = s5.shapes.add_textbox(Inches(7.6), Inches(4.6), Inches(4.8), Inches(2.2))
tf_r5_b = tx_r5_body.text_frame
tf_r5_b.word_wrap = True
r5_bullets = [
    "275 Backend Tests + 79 React Native Frontend Tests",
    "50-concurrent webhook race test verified against double-deposit",
    "Exact NUMERIC(15,2) decimal precision tested across all currencies",
    "End-to-end sandbox deployment verified with mock aggregators"
]
for b in r5_bullets:
    p = tf_r5_b.add_paragraph()
    p.space_before = Pt(10)
    p.text = f"✓  {b}"
    p.font.size = Pt(14)
    p.font.color.rgb = C_TEXT_MUTED

add_notes(s5, "Most fintech pitches at this stage have a Figma mockup. Croe has a fully implemented, tested backend and mobile app with 354 automated tests.")

# ==============================================================================
# SLIDE 6: Regulatory Strategy
# ==============================================================================
s6 = prs.slides.add_slide(blank_layout)
set_slide_background(s6, C_LIGHT_BG)
add_header(s6, "05 / Regulatory Strategy", "Escrow is a regulated activity, not a feature.")

tx_top = s6.shapes.add_textbox(Inches(1.2), Inches(2.6), Inches(10.5), Inches(0.8))
tx_top.text_frame.word_wrap = True
p = tx_top.text_frame.paragraphs[0]
p.text = "Payment Systems and Services Act, 2019 (Act 987)"
p.font.bold = True
p.font.size = Pt(16)
p.font.color.rgb = C_GREEN
p2 = tx_top.text_frame.add_paragraph()
p2.space_before = Pt(8)
p2.text = "Holding third-party funds in Ghana falls under Bank of Ghana regulation. Most escrow attempts fail either by ignoring the law and getting shut down, or by stalling for years trying to obtain an impossible licence."
p2.font.size = Pt(16)
p2.font.color.rgb = C_TEXT_BODY

# Column 1
tx_oth = s6.shapes.add_textbox(Inches(1.2), Inches(4.2), Inches(4.8), Inches(3.0))
tf_oth = tx_oth.text_frame
tf_oth.word_wrap = True
p = tf_oth.paragraphs[0]
p.text = "What Competitors Do (and Fail)"
p.font.bold = True
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(220, 38, 38)
for item in [
    "Operate unlicensed custody → Cease-and-desist from BoG",
    "Commingle customer float with operating capital",
    "Wait for full BoG licence → Millions in capital, multi-year stall",
    "Treat escrow as a pure UI feature on top of raw P2P transfers"
]:
    p = tf_oth.add_paragraph()
    p.space_before = Pt(12)
    p.text = f"✗  {item}"
    p.font.size = Pt(14)
    p.font.color.rgb = C_TEXT_BODY

# Column 2
tx_croe = s6.shapes.add_textbox(Inches(6.8), Inches(4.2), Inches(5.3), Inches(3.0))
tf_c = tx_croe.text_frame
tf_c.word_wrap = True
p = tf_c.paragraphs[0]
p.text = "The Croe Solution: Non-Custodial Tech"
p.font.bold = True
p.font.size = Pt(18)
p.font.color.rgb = C_BLUE
for item in [
    "Croe NEVER holds customer funds on its own corporate balance sheet",
    "Operates as a Technology Provider (PFTSP pathway under Act 987)",
    "A licensed entity holds the float (aggregator in pilot, bank trust account at scale)",
    "CustodyProvider abstraction allows swapping custody rails with zero code changes"
]:
    p = tf_c.add_paragraph()
    p.space_before = Pt(12)
    p.text = f"✓  {item}"
    p.font.size = Pt(14)
    p.font.color.rgb = C_TEXT_HEAD

add_notes(s6, "This slide decides the pitch. Investors in African fintech know Act 987. Answering the regulatory question upfront separates us from naive attempts.")

# ==============================================================================
# SLIDE 7: Phased Custody
# ==============================================================================
s7 = prs.slides.add_slide(blank_layout)
set_slide_background(s7, C_LIGHT_BG)
add_header(s7, "06 / Phased Custody", "Compliance designed in, not retrofitted.")

phases = [
    ("P0 · BUILD", "Sandbox / Dev", "Nobody (Mock data)", "Development & 354 automated tests complete"),
    ("P1 · PILOT", "Supervised Pilot", "Aggregator settlement account", "Low-volume supervised pilot; interim disclosure to users"),
    ("P2 · PARTNER", "Commercial Scale", "Licensed Bank Pooled Trust", "Fully compliant institutional launch; bank-held trust account"),
    ("P3 · LICENCE", "Own BoG Licence", "Croe (Dedicated DEMI/EPSP)", "Long-term future; only when processing volume justifies capital")
]

for idx, (phase, role, float_holder, status_desc) in enumerate(phases):
    top = Inches(2.8 + idx * 0.9)
    
    tx = s7.shapes.add_textbox(Inches(1.2), top, Inches(10.5), Inches(0.8))
    tf = tx.text_frame
    p = tf.paragraphs[0]
    p.text = f"{phase}   -   {role}"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = C_BLUE if idx == 2 else C_TEXT_HEAD
    
    p_sub = tf.add_paragraph()
    p_sub.text = f"Float Custody: {float_holder}   |   Status: {status_desc}"
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = C_TEXT_BODY

    # Separator
    if idx < 3:
        line = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), top + Inches(0.75), Inches(10.0), Pt(1))
        line.fill.solid()
        line.fill.fore_color.rgb = C_OFF_WHITE
        line.line.fill.background()

tx_n = s7.shapes.add_textbox(Inches(1.2), Inches(6.5), Inches(10.5), Inches(0.5))
p = tx_n.text_frame.paragraphs[0]
p.text = "Architectural Advantage: The CustodyProvider interface means moving from P1 to P2 is a configuration change, not a code rewrite."
p.font.bold = True
p.font.size = Pt(13)
p.font.color.rgb = C_GREEN

add_notes(s7, "The CustodyProvider abstraction means escrow logic never touches an aggregator SDK directly. A bank partner provides compliance and float holding.")

# ==============================================================================
# SLIDE 8: Disputes
# ==============================================================================
s8 = prs.slides.add_slide(blank_layout)
set_slide_background(s8, C_LIGHT_BG)
add_header(s8, "07 / Automated Arbitration", "Escrow only works if disputes resolve fairly and cheaply.")

dispute_tiers = [
    ("Tier 1 · Forensics", "Cryptographic Integrity", "SHA-256 media hashing catches recycled scam photos. Header capture (IP, device fingerprint, network) builds an immutable evidence chain.", C_BLUE),
    ("Tier 2 · Heuristics", "Sub-Second SQL Rules", "Deterministic SQL heuristics evaluate transaction velocity, Sybil ring detection, and structuring patterns in under 500 milliseconds.", RGBColor(220, 38, 38)),
    ("Tier 3 · AI Arbitrator", "Self-Hosted LLM", "Autonomous adjudication only at confidence >= 0.900. Truly ambiguous claims automatically escalate to Human L3 Review.", C_GREEN)
]

for idx, (title, sub, body, acc_col) in enumerate(dispute_tiers):
    left = Inches(1.2 + idx * 3.7)
    
    tx_p = s8.shapes.add_textbox(left, Inches(2.8), Inches(3.2), Inches(0.4))
    p = tx_p.text_frame.paragraphs[0]
    p.text = sub.upper()
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = acc_col
    
    tx = s8.shapes.add_textbox(left, Inches(3.2), Inches(3.3), Inches(2.6))
    tf = tx.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.bold = True
    p1.font.size = Pt(20)
    p1.font.color.rgb = C_TEXT_HEAD
    p2 = tf.add_paragraph()
    p2.space_before = Pt(16)
    p2.text = body
    p2.font.size = Pt(15)
    p2.font.color.rgb = C_TEXT_BODY

tx_pr = s8.shapes.add_textbox(Inches(1.2), Inches(6.2), Inches(10.5), Inches(0.6))
p = tx_pr.text_frame.paragraphs[0]
p.text = "Strict Data Privacy (Ghana Act 843): Customer dispute evidence NEVER leaves self-hosted infrastructure. No external third-party LLM APIs."
p.font.bold = True
p.font.size = Pt(13)
p.font.color.rgb = C_TEXT_MUTED

add_notes(s8, "Manual support destroys unit economics on small orders. Automating Tier 1 and 2 is what makes Croe profitable on GHS 1,200 orders.")

# ==============================================================================
# SLIDE 9: Business Model & Unit Economics
# ==============================================================================
s9 = prs.slides.add_slide(blank_layout)
set_slide_background(s9, C_OFF_WHITE)
add_header(s9, "08 / Business Model", "4% all-in: 2.5% vendor + 1.5% buyer. Zero fee on refunds.")

tx_lead = s9.shapes.add_textbox(Inches(1.2), Inches(2.8), Inches(4.5), Inches(4.3))
tf_l = tx_lead.text_frame
tf_l.word_wrap = True

p = tf_l.paragraphs[0]
p.text = "LEAD VERTICAL: PHONES, ELECTRONICS & SNEAKERS"
p.font.bold = True
p.font.size = Pt(12)
p.font.color.rgb = C_GREEN

p2 = tf_l.add_paragraph()
p2.space_before = Pt(10)
p2.text = "GHS 1,200 Average Order Value"
p2.font.bold = True
p2.font.size = Pt(22)
p2.font.color.rgb = C_TEXT_HEAD

lead_stats = [
    ("Croe Fee (4%)", "GHS 48.00"),
    ("MoMo Collection (1.95%)", "- GHS 23.75"),
    ("Payout, ID Check, SMS, Hosting", "- GHS 2.55"),
    ("Fraud Reserve & Refunds", "- GHS 4.90"),
    ("Croe Net Margin / Order", "GHS 16.80"),
    ("Company Break-Even", "~37,900 orders / mo")
]
for label, val in lead_stats:
    p = tf_l.add_paragraph()
    p.space_before = Pt(12)
    p.text = f"{label}:  {val}"
    p.font.size = Pt(15)
    p.font.color.rgb = C_TEXT_BODY
    if "Net Margin" in label or "Break-Even" in label:
        p.font.bold = True
        p.font.color.rgb = C_TEXT_HEAD

line = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.3), Inches(2.8), Pt(1), Inches(3.5))
line.fill.solid()
line.fill.fore_color.rgb = C_LINE
line.line.fill.background()

tx_base = s9.shapes.add_textbox(Inches(7.2), Inches(2.8), Inches(4.5), Inches(4.3))
tf_b = tx_base.text_frame
tf_b.word_wrap = True

p = tf_b.paragraphs[0]
p.text = "BASE CASE: BLENDED ORDER"
p.font.bold = True
p.font.size = Pt(12)
p.font.color.rgb = C_BLUE

p2 = tf_b.add_paragraph()
p2.space_before = Pt(10)
p2.text = "GHS 800 Average Order Value"
p2.font.bold = True
p2.font.size = Pt(22)
p2.font.color.rgb = C_TEXT_HEAD

base_stats = [
    ("Croe Fee (4%)", "GHS 32.00"),
    ("MoMo Collection (1.95%)", "- GHS 15.83"),
    ("Payout, ID Check, SMS, Hosting", "- GHS 2.55"),
    ("Fraud Reserve & Refunds", "- GHS 3.29"),
    ("Croe Net Margin / Order", "GHS 10.33"),
    ("Company Break-Even", "~61,700 orders / mo")
]
for label, val in base_stats:
    p = tf_b.add_paragraph()
    p.space_before = Pt(12)
    p.text = f"{label}:  {val}"
    p.font.size = Pt(15)
    p.font.color.rgb = C_TEXT_BODY
    if "Net Margin" in label or "Break-Even" in label:
        p.font.bold = True
        p.font.color.rgb = C_TEXT_HEAD

add_notes(s9, "Say the uncomfortable part first: MoMo collection costs 1.95% and each payout GHS 1.00, so a 2% fee would lose GHS 4-7 on every order. At 4% all-in (2.5% vendor + 1.5% buyer protection fee) net is 1.29% on a GHS 800 order. Against a fully staffed month-36 cost base of GHS 637k/month, company break-even is ~62,000 orders/month. Levers: a 5% fee (~34,800), direct telco rails via our own PSP licence (~34,500), high-ticket vendors (~37,900 at GHS 1,200). Buyer acceptance of the 1.5% fee is the first thing the pilot tests.")

# ==============================================================================
# SLIDE 10: Go To Market
# ==============================================================================
s10 = prs.slides.add_slide(blank_layout)
set_slide_background(s10, C_LIGHT_BG)
add_header(s10, "09 / Go To Market", "Vendors bring their own buyers (B2B2C).")

segments = [
    ("Phones & Electronics", "GHS 1k – 3.5k", "Highest fraud anxiety. Buyers routinely ghost at 'send MoMo first'. Vendors desperate for guaranteed closing.", C_BLUE),
    ("Sneakers & Streetwear", "GHS 600 – 1.8k", "Instagram-native sneaker resellers. High dispute rates over authenticity. Perfect for escrow inspection window.", C_GREEN),
    ("Curated Thrift", "GHS 400 – 900", "High velocity, intercity dispatch. Vendors lose repeat orders when distant buyers refuse to pay before courier delivery.", C_TEXT_MUTED)
]

for idx, (title, sub, desc, col) in enumerate(segments):
    left = Inches(1.2 + idx * 3.7)
    
    tx = s10.shapes.add_textbox(left, Inches(2.8), Inches(3.2), Inches(3.0))
    tf = tx.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.bold = True
    p1.font.size = Pt(20)
    p1.font.color.rgb = C_TEXT_HEAD
    p_sub = tf.add_paragraph()
    p_sub.space_before = Pt(8)
    p_sub.text = sub
    p_sub.font.bold = True
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = col
    p2 = tf.add_paragraph()
    p2.space_before = Pt(16)
    p2.text = desc
    p2.font.size = Pt(15)
    p2.font.color.rgb = C_TEXT_BODY

tx_pb = s10.shapes.add_textbox(Inches(1.2), Inches(5.8), Inches(10.5), Inches(0.8))
tf_pb = tx_pb.text_frame
tf_pb.word_wrap = True
p = tf_pb.paragraphs[0]
p.text = "THE PITCH TO A VENDOR IS CONVERSION, NOT SECURITY:"
p.font.bold = True
p.font.size = Pt(12)
p.font.color.rgb = C_GREEN
p2 = tf_pb.add_paragraph()
p2.text = "\"You will close the 25% of qualified buyers who currently ghost your MoMo details.\""
p2.font.name = FONT_TITLE
p2.font.bold = True
p2.font.size = Pt(22)
p2.font.color.rgb = C_TEXT_HEAD

add_notes(s10, "Buyers won't search an app store for an escrow app. Vendors recruit buyers because Croe closes sales they are currently losing.")

# ==============================================================================
# SLIDE 11: Traction & Validation Plan
# ==============================================================================
s11 = prs.slides.add_slide(blank_layout)
set_slide_background(s11, C_LIGHT_BG)
add_header(s11, "10 / Pilot Metrics", "Pre-launch. What we are validating, and when.")

tx_c = s11.shapes.add_textbox(Inches(1.2), Inches(2.8), Inches(4.5), Inches(4.2))
tf_c = tx_c.text_frame
tf_c.word_wrap = True
p = tf_c.paragraphs[0]
p.text = "Current Position"
p.font.bold = True
p.font.size = Pt(20)
p.font.color.rgb = C_BLUE

cur_items = [
    "Product complete and tested (354 tests passing)",
    "Running end-to-end on aggregator sandbox",
    "10–20 pilot vendors in recruitment cohort in Accra",
    "Entity formation & banking partner discussions active",
    "Pre-revenue by design: verifying compliance before live money"
]
for item in cur_items:
    p = tf_c.add_paragraph()
    p.space_before = Pt(14)
    p.text = f"•  {item}"
    p.font.size = Pt(14)
    p.font.color.rgb = C_TEXT_BODY

line = s11.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.3), Inches(2.8), Pt(1), Inches(3.5))
line.fill.solid()
line.fill.fore_color.rgb = C_LINE
line.line.fill.background()

tx_g = s11.shapes.add_textbox(Inches(7.2), Inches(2.8), Inches(5.0), Inches(4.2))
tf_g = tx_g.text_frame
tf_g.word_wrap = True
p = tf_g.paragraphs[0]
p.text = "Stage 1 Pilot Gates"
p.font.bold = True
p.font.size = Pt(20)
p.font.color.rgb = C_TEXT_HEAD

gates = [
    ("Volume", "≥ 250 orders by month 6", "Real demand on capped live rails"),
    ("Dispute Rate", "< 8% of transactions", "Validates fraud resistance"),
    ("AI Accuracy", "> 95% upheld", "Proves 3-tier dispute engine"),
    ("Buyer Drop-off", "< 30% with 1.5% fee shown", "Validates 4% pricing and web/USSD UX"),
    ("Retention", "> 60% active", "Proves commercial utility"),
    ("Reconciliation", "0 mismatches", "Mandatory requirement")
]
for metric, tgt, why in gates:
    p = tf_g.add_paragraph()
    p.space_before = Pt(12)
    p.text = f"{metric}: {tgt}"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = C_TEXT_HEAD
    p2 = tf_g.add_paragraph()
    p2.text = f"  {why}"
    p2.font.size = Pt(14)
    p2.font.color.rgb = C_TEXT_BODY

add_notes(s11, "Be straightforwardly honest: we are pre-revenue. In fintech, one invented number ends the conversation. We know the exact numbers that prove or kill this business.")

# ==============================================================================
# SLIDE 12: The Ask
# ==============================================================================
s12 = prs.slides.add_slide(blank_layout)
set_slide_background(s12, C_OFF_WHITE)
add_header(s12, "11 / Capital Ask", "US$387k pre-seed to a licensed, partner-held launch.")

tx_ta = s12.shapes.add_textbox(Inches(1.2), Inches(2.8), Inches(5.0), Inches(3.8))
tf_ta = tx_ta.text_frame
tf_ta.word_wrap = True
p = tf_ta.paragraphs[0]
p.text = "STAGE 1: LEGAL PILOT (MONTHS 1–6)"
p.font.bold = True
p.font.size = Pt(12)
p.font.color.rgb = C_BLUE
p2 = tf_ta.add_paragraph()
p2.space_before = Pt(8)
p2.text = "US$26k  (GHS 305k)"
p2.font.bold = True
p2.font.size = Pt(26)
p2.font.color.rgb = C_TEXT_HEAD

ta_items = [
    "Company, TIN, DPC registration, lawyer's custody opinion",
    "Written aggregator approval; BoG sandbox or partner LOI",
    "Founder stipends, part-time ops, lean hosting (AI on CPU)",
    "10–50 vendors; ~500 capped, disclosed live orders",
    "Gate: buyer fee acceptance, order value, dispute rate"
]
for item in ta_items:
    p = tf_ta.add_paragraph()
    p.space_before = Pt(12)
    p.text = f"✓  {item}"
    p.font.size = Pt(14)
    p.font.color.rgb = C_TEXT_BODY

line = s12.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.5), Inches(2.8), Pt(1), Inches(3.5))
line.fill.solid()
line.fill.fore_color.rgb = C_LINE
line.line.fill.background()

tx_tb = s12.shapes.add_textbox(Inches(7.2), Inches(2.8), Inches(5.0), Inches(3.8))
tf_tb = tx_tb.text_frame
tf_tb.word_wrap = True
p = tf_tb.paragraphs[0]
p.text = "STAGE 2: COMPLIANT LAUNCH (MONTHS 7–18)"
p.font.bold = True
p.font.size = Pt(12)
p.font.color.rgb = C_GREEN
p2 = tf_tb.add_paragraph()
p2.space_before = Pt(8)
p2.text = "US$360k  (GHS 4.18M)"
p2.font.bold = True
p2.font.size = Pt(26)
p2.font.color.rgb = C_TEXT_HEAD

tb_items = [
    "BoG PFTSP licence + partner-bank trust account (P2)",
    "The five BoG management roles; ~12 staff",
    "External pentest, ISO 27001, full BoG policy set",
    "Float (one day of order value) + partner reserve",
    "Milestone: ~5,000 orders/month on compliant rails"
]
for item in tb_items:
    p = tf_tb.add_paragraph()
    p.space_before = Pt(12)
    p.text = f"✓  {item}"
    p.font.size = Pt(14)
    p.font.color.rgb = C_TEXT_BODY

add_notes(s12, "Tie the number to the milestone. Pre-seed US$387k (range US$189k-713k) covers stages 1-2 over 18 months, with 15% contingency and working capital; FX US$1 = GHS 11.60. Most of stage 2 is people: the BoG names five management roles in its licensing pack. A seed round of ~US$916k around month 15 funds scale to ~30,000 orders/month; US$1.30M total over 36 months. Line items and sources: FINANCIAL-MODEL.md.")

# ==============================================================================
# SLIDE 13: Appendix Divider (Dark Theme)
# ==============================================================================
s13 = prs.slides.add_slide(blank_layout)
set_slide_background(s13, C_DARK_BG)

logo_path = os.path.join(ASSETS_DIR, "croe_wordmark_light.png")
if os.path.exists(logo_path):
    s13.shapes.add_picture(logo_path, Inches(1.2), Inches(2.2), width=Inches(3.5))

tx_app = s13.shapes.add_textbox(Inches(1.15), Inches(3.4), Inches(10.5), Inches(2.5))
tf_app = tx_app.text_frame
p = tf_app.paragraphs[0]
p.text = "Appendix"
p.font.name = FONT_TITLE
p.font.size = Pt(56)
p.font.bold = True
p.font.color.rgb = C_WHITE

p2 = tf_app.add_paragraph()
p2.space_before = Pt(20)
p2.text = "Technical Depth, Risk Registry, Expansion, and Diligence FAQ"
p2.font.size = Pt(22)
p2.font.color.rgb = RGBColor(148, 163, 184)

add_notes(s13, "Appendix slides are for Q&A, not for presenting.")

# ==============================================================================
# SLIDE 14: Appendix A1 - Why a Student Built This
# ==============================================================================
s14 = prs.slides.add_slide(blank_layout)
set_slide_background(s14, C_LIGHT_BG)
add_header(s14, "Appendix A1", "Why a student built this: Discipline & Engineering Rules.")

tx_a1 = s14.shapes.add_textbox(Inches(1.2), Inches(2.8), Inches(10.5), Inches(4.2))
tf_a1 = tx_a1.text_frame
tf_a1.word_wrap = True

p = tf_a1.paragraphs[0]
p.text = "The Specification Package is the Due-Diligence Pack"
p.font.bold = True
p.font.size = Pt(20)
p.font.color.rgb = C_BLUE

a1_points = [
    "28 comprehensive specification documents written before a single line of application code was authored.",
    "Strict Phase-Gate Rule (PROC-01): No phase begins until the previous phase is production-grade, complete, and tested.",
    "Mandatory engineering constraints: NUMERIC(15,2) decimal precision, append-only sub-ledger, SELECT FOR UPDATE locking.",
    "Zero vanity architecture: Self-hosted open-weights LLM running locally; P0 build cost literally ≈ GHS 0.00.",
    "Full due-diligence package available immediately: Architecture, Threat Model, DDL schema, and 354 passing tests."
]
for pt in a1_points:
    p = tf_a1.add_paragraph()
    p.space_before = Pt(14)
    p.text = f"•  {pt}"
    p.font.size = Pt(15)
    p.font.color.rgb = C_TEXT_BODY

add_notes(s14, "If asked how a student built this: answer with the 28-document specification package and phase-gate discipline. Offer the spec package as the diligence pack.")

# ==============================================================================
# SLIDE 15: Appendix A2 - Technical Depth
# ==============================================================================
s15 = prs.slides.add_slide(blank_layout)
set_slide_background(s15, C_LIGHT_BG)
add_header(s15, "Appendix A2", "Financial Concurrency & Infrastructure Depth.")

cards_a2 = [
    ("Double-Spend Defense", "3-Layer Gate", "Layer 1: Redis SETNX fast dedup (24h TTL).\nLayer 2: SELECT FOR UPDATE pessimistic row locks.\nLayer 3: Partial unique DB indexes (Error 23505 trap)."),
    ("Ordering Rule", "Pay Then Ledger", "Negative ledger entries are recorded ONLY after the external provider confirms disbursement success. Failed payouts leave balance intact and retriable."),
    ("Reconciliation", "3-Way Daily Match", "Automated daily match: Provider Statement vs Sub-Ledger Sum vs Rail Log. Any non-zero delta immediately freezes disbursements and pages on-call.")
]
for idx, (title, sub, body) in enumerate(cards_a2):
    left = Inches(1.2 + idx * 3.7)
    
    tx = s15.shapes.add_textbox(left, Inches(2.8), Inches(3.2), Inches(4.2))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = Pt(20)
    p.font.color.rgb = C_BLUE
    p_sub = tf.add_paragraph()
    p_sub.space_before = Pt(8)
    p_sub.text = sub.upper()
    p_sub.font.bold = True
    p_sub.font.size = Pt(12)
    p_sub.font.color.rgb = C_GREEN
    p2 = tf.add_paragraph()
    p2.space_before = Pt(16)
    p2.text = body
    p2.font.size = Pt(14)
    p2.font.color.rgb = C_TEXT_BODY

add_notes(s15, "Highlight that financial integrity is handled at the database and architecture level, not assumed in application memory.")

# ==============================================================================
# SLIDE 16: Appendix A3 - Risks, Stated Plainly
# ==============================================================================
s16 = prs.slides.add_slide(blank_layout)
set_slide_background(s16, C_OFF_WHITE)
add_header(s16, "Appendix A3", "Key Business Risks & Concrete Mitigations.")

risks = [
    ("Aggregator restricts escrow model", "Clarified in writing during Week 1. If restricted, we jump directly to licensed P2 partner bank."),
    ("P1 is legally thin at scale", "Deliberately an interim state. Volumes strictly capped; explicit disclosure; active sprint toward P2 trust account."),
    ("Working capital gap (T+2 settlement)", "Float sized and reserved in budget; option to align vendor payout timing to settlement cycle."),
    ("Buyer drop-off at install step", "Web-based payment fallback supported; buyer pays via mobile browser + USSD prompt without app install."),
    ("Dispute volume overwhelms team", "Automated Tier 1 (media hash) and Tier 2 (SQL heuristics) filter out clear fraud before human review."),
    ("Wrong AI dispute decision", "Written adjudication policy, ring-fenced loss reserve plus 0.3% of order value per order, appeal path, and ToS liability caps.")
]
for idx, (risk, mit) in enumerate(risks):
    top = Inches(2.8 + idx * 0.7)
    tx = s16.shapes.add_textbox(Inches(1.2), top, Inches(10.5), Inches(0.5))
    tf = tx.text_frame
    p = tf.paragraphs[0]
    p.text = f"Risk: {risk}"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(220, 38, 38)
    p2 = tf.add_paragraph()
    p2.text = f"Mitigation: {mit}"
    p2.font.size = Pt(14)
    p2.font.color.rgb = C_TEXT_BODY

add_notes(s16, "Volunteering risks with concrete mitigations signals competence and maturity.")

# ==============================================================================
# SLIDE 17: Appendix A4 - Regional Expansion
# ==============================================================================
s17 = prs.slides.add_slide(blank_layout)
set_slide_background(s17, C_LIGHT_BG)
add_header(s17, "Appendix A4", "Modular Regional Expansion: Kenya & Nigeria.")

exp_cards = [
    ("Ghana (Current)", "GHS · MTN MoMo, Telecel", "Bank of Ghana (Act 987). Lead launch market. 20M+ wallets. Initial P1 aggregator pilot → P2 partner trust account."),
    ("Kenya (Fast Follow)", "KES · Safaricom M-Pesa", "Central Bank of Kenya. Dominant M-Pesa rail (30M+ users). Added via Safaricom Daraja API PaymentRail implementation."),
    ("Nigeria (High Volume)", "NGN · Bank Transfer, Cards", "Central Bank of Nigeria. Card networks, NIP bank transfers, and BVN identity verification. High social commerce volume.")
]
for idx, (title, rail, desc) in enumerate(exp_cards):
    left = Inches(1.2 + idx * 3.7)
    
    tx = s17.shapes.add_textbox(left, Inches(2.8), Inches(3.2), Inches(4.2))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = Pt(20)
    p.font.color.rgb = C_BLUE
    p_sub = tf.add_paragraph()
    p_sub.space_before = Pt(8)
    p_sub.text = rail
    p_sub.font.bold = True
    p_sub.font.size = Pt(12)
    p_sub.font.color.rgb = C_GREEN
    p2 = tf.add_paragraph()
    p2.space_before = Pt(16)
    p2.text = desc
    p2.font.size = Pt(14)
    p2.font.color.rgb = C_TEXT_BODY

add_notes(s17, "Currency and rails sit behind clean abstractions. Kenya and Nigeria are modular PaymentRail additions, not codebase rewrites.")

# ==============================================================================
# SLIDE 18: Appendix A5 - Frequently Asked Questions
# ==============================================================================
s18 = prs.slides.add_slide(blank_layout)
set_slide_background(s18, C_LIGHT_BG)
add_header(s18, "Appendix A5", "Investor & Partner FAQ.")

faqs = [
    ("Isn't this unlicensed money transmission?", "No. Croe never holds client funds on its own balance sheet. In pilot, funds settle via aggregator merchant settlement; at scale, via a licensed partner bank trust account."),
    ("What stops the buyer confirming delivery and then lying?", "Dispute window closes upon buyer confirmation. Inspection timers auto-release funds. Media is SHA-256 hashed to prevent recycled scam photos."),
    ("Why would a vendor wait for payout?", "Because they currently lose 25% of sales to payment anxiety. Closing 8 additional GHS 1,200 sales per month far outweighs waiting 24 hours for settlement."),
    ("Why won't MTN or Paystack just copy this?", "Aggregators provide generic rails; they avoid subjective dispute arbitration and merchant operational friction. Our moat is vendor trust and dispute fairness.")
]
for idx, (q, a) in enumerate(faqs):
    top = Inches(2.8 + idx * 1.0)
    
    tx = s18.shapes.add_textbox(Inches(1.2), top, Inches(10.5), Inches(0.8))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"Q: {q}"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = C_BLUE
    p2 = tf.add_paragraph()
    p2.space_before = Pt(6)
    p2.text = f"A: {a}"
    p2.font.size = Pt(14)
    p2.font.color.rgb = C_TEXT_BODY

add_notes(s18, "Keep answers crisp, confident, and grounded in architecture.")

try:
    prs.save(OUTPUT_PPTX)
    print(f"Presentation generated successfully: {OUTPUT_PPTX}")
except PermissionError:
    print(f"Permission denied: {OUTPUT_PPTX}")
    print("Please close PowerPoint and run this script again.")
