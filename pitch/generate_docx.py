#!/usr/bin/env python3
"""
Generate Croe Stakeholder Word Documents (.docx)
Creates executive-styled Word documents adhering to the Croe Design System ("Calm over confrontation",
trademark 'croe.' with green dot, clean typography, styled tables, and callouts).
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

ASSETS_DIR = r"c:\Users\Jerry\Desktop\PROJECT 2026\Croe\pitch\assets"
PITCH_DIR = r"c:\Users\Jerry\Desktop\PROJECT 2026\Croe\pitch"

# Color constants
COLOR_NAVY_HEX = "0F172A"       # Slate 900
COLOR_BLUE_HEX = "2563EB"       # Trust Blue
COLOR_GREEN_HEX = "1FC16B"      # Croe Green
COLOR_LIGHT_BG_HEX = "F8FAFC"   # Slate 50
COLOR_BORDER_HEX = "E2E8F0"     # Border
COLOR_BLUE_BG_HEX = "EFF6FF"    # Light Blue Tint
COLOR_AMBER_BG_HEX = "FEF3C7"   # Warm Amber Tint
COLOR_AMBER_TEXT_HEX = "B45309" # Dark Amber

RGB_HEAD = RGBColor(15, 23, 42)
RGB_BODY = RGBColor(51, 65, 85)
RGB_MUTED = RGBColor(100, 116, 139)
RGB_BLUE = RGBColor(37, 99, 235)
RGB_GREEN = RGBColor(31, 193, 107)

FONT_PRIMARY = "Segoe UI"

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'))

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def style_table(table, header_bg=COLOR_NAVY_HEX, header_text_color=RGBColor(255, 255, 255), zebra=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Set thin border
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_HEX}"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{COLOR_BORDER_HEX}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_HEX}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

    # Style rows
    for row_idx, row in enumerate(table.rows):
        is_header = (row_idx == 0)
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=130, bottom=130, left=160, right=160)
            if is_header:
                set_cell_shading(cell, header_bg)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs:
                        run.font.name = FONT_PRIMARY
                        run.font.bold = True
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = header_text_color
            else:
                if zebra and row_idx % 2 == 1:
                    set_cell_shading(cell, COLOR_LIGHT_BG_HEX)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = FONT_PRIMARY
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = RGB_BODY

def add_header_banner(doc, title, subtitle, doc_ref):
    # Top table with logo on left and doc ref on right
    hdr_table = doc.add_table(rows=1, cols=2)
    hdr_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_table.autofit = False
    hdr_table.columns[0].width = Inches(4.5)
    hdr_table.columns[1].width = Inches(2.5)

    c0 = hdr_table.cell(0, 0)
    c1 = hdr_table.cell(0, 1)

    # Logo
    logo_path = os.path.join(ASSETS_DIR, "croe_wordmark_dark.png")
    p0 = c0.paragraphs[0]
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(0)
    if os.path.exists(logo_path):
        p0.add_run().add_picture(logo_path, width=Inches(1.2))

    # Ref / Status
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.space_before = Pt(4)
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run(doc_ref)
    r1.font.name = FONT_PRIMARY
    r1.font.size = Pt(9)
    r1.font.bold = True
    r1.font.color.rgb = RGB_BLUE

    # Title & Subtitle
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_before = Pt(16)
    p_t.paragraph_format.space_after = Pt(4)
    r_t = p_t.add_run(title)
    r_t.font.name = FONT_PRIMARY
    r_t.font.size = Pt(22)
    r_t.font.bold = True
    r_t.font.color.rgb = RGB_HEAD

    p_s = doc.add_paragraph()
    p_s.paragraph_format.space_before = Pt(0)
    p_s.paragraph_format.space_after = Pt(16)
    r_s = p_s.add_run(subtitle)
    r_s.font.name = FONT_PRIMARY
    r_s.font.size = Pt(12)
    r_s.font.color.rgb = RGB_MUTED

    # Subtle divider
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(14)
    r_div = p_div.add_run("―" * 58)
    r_div.font.size = Pt(8)
    r_div.font.color.rgb = RGBColor(226, 232, 240)

def add_heading_styled(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        r.font.name = FONT_PRIMARY
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGB_BLUE
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.name = FONT_PRIMARY
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = RGB_HEAD
    return p

def add_body_p(doc, text, bold_prefix="", space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = FONT_PRIMARY
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGB_HEAD
    r = p.add_run(text)
    r.font.name = FONT_PRIMARY
    r.font.size = Pt(10)
    r.font.color.rgb = RGB_BODY
    return p

def add_callout(doc, text, title="KEY STRATEGIC INSIGHT", bg_hex=COLOR_BLUE_BG_HEX, border_hex=COLOR_BLUE_HEX, title_rgb=RGB_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(7.0)
    cell = table.cell(0, 0)
    set_cell_shading(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    # Left border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    r_t = p.add_run(f"⚡  {title}")
    r_t.font.name = FONT_PRIMARY
    r_t.font.size = Pt(9.5)
    r_t.font.bold = True
    r_t.font.color.rgb = title_rgb

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(text)
    r2.font.name = FONT_PRIMARY
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGB_HEAD

    # Space after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)

def add_footer(doc, doc_title):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(f"{doc_title}  |  croe.  |  Confidential")
    r.font.name = FONT_PRIMARY
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGB_MUTED

def fill_table(doc, headers, rows, header_bg=COLOR_NAVY_HEX, zebra=True):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    for i, h in enumerate(headers):
        table.cell(0, i).paragraphs[0].text = h
    for r_idx, r_data in enumerate(rows, start=1):
        for c_idx, val in enumerate(r_data):
            table.cell(r_idx, c_idx).paragraphs[0].text = val
    style_table(table, header_bg=header_bg, zebra=zebra)
    return table

# ==============================================================================
# DOC 1: Executive One-Pager (.docx)
# ==============================================================================
def generate_one_pager_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)

    add_header_banner(
        doc,
        title="croe. - Executive Summary",
        subtitle="Escrow for Social Commerce, on Mobile Money. Ghana-First.",
        doc_ref="STAKEHOLDER BRIEF · SEPTEMBER 2026"
    )

    add_heading_styled(doc, "1. The Problem: The Trust Standoff in Social Commerce", level=1)
    add_body_p(doc, "Social commerce in Ghana runs on trust that does not exist. A buyer finds a smartphone, laptop, or designer sneakers on Instagram and is asked to send Mobile Money (MoMo) to a stranger. The vendor is asked to dispatch goods via dispatch rider to someone who may refuse to pay. Both positions are completely rational, and the sale dies between them.", bold_prefix="Structural Failure: ")
    add_body_p(doc, "Vendors routinely lose 25%–30% of qualified buyer inquiries to payment anxiety. Buyers who do pay upfront absorb recurring fraud with zero recourse - no enforceable receipt, no neutral arbiter, and no path to an automated refund.")

    add_heading_styled(doc, "2. The Solution: Frictionless Escrow on MoMo Rails", level=1)
    
    # 4-step table
    t_sol = doc.add_table(rows=5, cols=3)
    headers = ["Step", "Action", "Customer Experience"]
    for i, h in enumerate(headers):
        t_sol.cell(0, i).paragraphs[0].text = h
    
    data_sol = [
        ("1. Link", "Vendor creates escrow link", "Shared directly in WhatsApp or Instagram DM where the sale is already happening."),
        ("2. Pay", "Buyer authorizes MoMo payment", "Standard USSD prompt on their handset. Funds lock in escrow - vendor sees it, but cannot touch it."),
        ("3. Ship", "Vendor dispatches order", "Goods dispatched against guaranteed locked money rather than an empty promise."),
        ("4. Release", "Buyer confirms delivery", "Vendor paid automatically, less 2.5% commission. If an issue occurs, a structured dispute opens.")
    ]
    for row_idx, (c0, c1, c2) in enumerate(data_sol, start=1):
        t_sol.cell(row_idx, 0).paragraphs[0].text = c0
        t_sol.cell(row_idx, 1).paragraphs[0].text = c1
        t_sol.cell(row_idx, 2).paragraphs[0].text = c2
    style_table(t_sol)

    add_heading_styled(doc, "3. Why Escrow Is Hard - And Why Croe Is Ahead", level=1)
    add_body_p(doc, "Holding third-party funds falls under the Bank of Ghana (BoG) and the Payment Systems and Services Act, 2019 (Act 987). Unlicensed platforms get shut down; full licences require millions in capital. Croe solves this through a Phased Custody Chain: Croe never holds client funds on its own balance sheet. In pilot (P1), funds settle via aggregator settlement; at scale (P2), via a licensed partner bank trust account. The CustodyProvider software abstraction means moving from pilot to institutional custody is a configuration switch, not a code rewrite.", bold_prefix="1. Regulatory Compliance by Design: ")
    add_body_p(doc, "Manual arbitration destroys unit economics on small orders. Croe runs a 3-tier dispute engine: cryptographic SHA-256 media hashing (catches recycled scam imagery across vendors), deterministic SQL heuristics (velocity and Sybil tripwires), and a self-hosted open-weights LLM arbitrator that handles straightforward claims and escalates only ambiguous cases to human review. Customer dispute data never leaves self-hosted infrastructure.", bold_prefix="2. Automated Impartial Dispute Triage: ")

    add_heading_styled(doc, "4. Current Status: Built, Tested & Proven", level=1)
    add_body_p(doc, "Croe is not a mockup. The complete backend and React Native mobile application are finished, running on aggregator sandboxes with 354 automated tests passing (275 backend + 79 frontend), including concurrent-webhook race tests and NUMERIC(15,2) decimal precision verification.")

    add_heading_styled(doc, "5. Business Model & Unit Economics", level=1)
    add_body_p(doc, "Croe earns 4.0% all-in on successfully released orders: a 2.5% vendor commission deducted from payout plus a 1.5% buyer protection fee shown at checkout. No fees on refunds. MoMo collection costs 1.95% and each payout GHS 1.00, so a 2% fee would lose money on every order.")

    fill_table(doc, ["Commercial Segment", "Avg Order", "Croe Fee (4%)", "Net / Order", "Orders/mo to Break-Even"], [
        ("Lead: Phones, Electronics & Sneakers", "GHS 1,200.00", "GHS 48.00", "GHS 16.80", "~37,900"),
        ("Base case (blended)", "GHS 800.00", "GHS 32.00", "GHS 10.33", "~61,700"),
        ("General Social Commerce", "GHS 450.00", "GHS 18.00", "GHS 4.67", "~136,300"),
    ], header_bg=COLOR_NAVY_HEX)

    add_callout(
        doc,
        "Net is 1.29% of order value on a GHS 800 order, after collection, payout, fraud reserve, ID checks, refunds, SMS and hosting. Against a fully staffed month-36 cost base of GHS 637,000/month, company break-even is ~62,000 orders/month. A 5% fee or direct telco rails cut that to ~35,000; high-ticket vendors cut it by ~40%.",
        title="PAYMENTS IS A VOLUME BUSINESS"
    )

    add_heading_styled(doc, "6. The Ask: Pre-Seed to a Licensed Launch", level=1)
    add_body_p(doc, "US$387k pre-seed (range US$189k-713k; GHS 4.48M base) takes Croe from sandbox to a licensed, partner-held launch in 18 months. Includes 15% contingency and working capital. FX: US$1 = GHS 11.60.")
    fill_table(doc, ["Stage", "Months", "Net Need (base)", "What It Buys", "Gate to Next Stage"], [
        ("1 · Legal pilot", "1-6", "US$26k", "Company, DPC registration, lawyer's custody opinion, written aggregator approval, BoG sandbox or partner LOI, 10-50 vendors, ~500 capped live orders.", "Written aggregator/partner yes; measured buyer drop-off, order value, dispute rate."),
        ("2 · Compliant launch", "7-18", "US$360k", "BoG PFTSP licence, partner-bank trust account, BoG management team (~12 staff), pentest, ISO 27001, Accra + Kumasi, ~5,000 orders/month.", "Clean daily reconciliation, disputes < 8%, vendor retention."),
    ], header_bg=COLOR_BLUE_HEX)
    add_body_p(doc, "A seed round of ~US$916k around month 15 funds stage 3: ~25 staff and ~30,000 orders/month by month 36. Total over 36 months: US$1.30M.", space_after=4)

    add_footer(doc, "Executive Summary")
    out_file = os.path.join(PITCH_DIR, "Croe_Executive_One_Pager.docx")
    doc.save(out_file)
    print(f"Generated: {out_file}")

# ==============================================================================
# DOC 2: Financial Model & Unit Economics (.docx)
# ==============================================================================
def generate_financial_model_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

    add_header_banner(
        doc,
        title="croe. - Financial Model & Unit Economics",
        subtitle="Per-Order Economics, Staged Capital Plan & Break-Even Sensitivity",
        doc_ref="FINANCIAL MODEL · 25 SEP 2026 · RULE FIN-01"
    )

    add_callout(
        doc,
        "The earlier model assumed a 2.0% vendor commission netting GHS 22.50 per order, fixed costs under GHS 500/month and break-even at 22-63 orders/month. It left out the 1.95% MoMo collection fee and every salary. This model is rebuilt from published 2026 prices. FX: US$1 = GHS 11.60. All capital figures include 15% contingency.",
        title="THIS SUPERSEDES THE EARLIER MODEL", bg_hex=COLOR_AMBER_BG_HEX, border_hex=COLOR_AMBER_TEXT_HEX, title_rgb=RGBColor(180, 83, 9)
    )

    add_heading_styled(doc, "1. Executive Summary of Numbers", level=1)
    fill_table(doc, ["Metric", "Value (base case)", "Range"], [
        ("Pricing", "4.0% all-in = 2.5% vendor commission + 1.5% buyer protection fee", "3.5% is the floor; 5% halves break-even"),
        ("Net per released order (GHS 800 order)", "GHS 10.33 (1.29% of order value)", "GHS 4.67 at GHS 450 · GHS 16.80 at GHS 1,200"),
        ("Capital to a licensed, compliant launch (stages 1-2, 18 months)", "US$387k (GHS 4.48M)", "US$189k - US$713k"),
        ("Capital to scale (stage 3, months 19-36)", "US$916k (GHS 10.63M)", "US$442k - US$1.80M"),
        ("Total over 36 months", "US$1.30M (GHS 15.11M)", "US$631k - US$2.52M"),
        ("Fixed cost at month 36", "GHS 637,000/month (US$54,900)", "Team of ~25"),
        ("Company break-even", "~62,000 released orders/month", "~34,800 at a 5% fee · ~34,500 with direct telco rails"),
    ], header_bg=COLOR_NAVY_HEX)
    add_body_p(doc, "Low pairs bottom-of-range costs with base-case volume. High pairs top-of-range costs with half the base-case volume.", space_after=4)

    add_heading_styled(doc, "2. Pricing: Why 2% Cannot Work", level=1)
    add_body_p(doc, "MoMo collection costs 1.95% of the amount collected (Paystack; Hubtel 1.95%, min GHS 0.30), and each MoMo payout costs GHS 1.00 (Paystack). At a 2.0% vendor-only commission the collection fee alone uses up 98% of the fee, and Croe loses GHS 4.19-6.83 on every order.", bold_prefix="Published fees: ")
    fill_table(doc, ["Party", "Fee", "On a GHS 1,200 order"], [
        ("Buyer", "1.5% protection fee, shown at checkout", "Pays GHS 1,218.00"),
        ("Vendor", "2.5% commission, deducted from payout", "Receives GHS 1,170.00"),
        ("Croe", "4.0% gross", "Earns GHS 48.00 gross -> GHS 16.80 net"),
        ("Refunded order", "No fee to either party", "Croe earns GHS 0.00"),
    ], header_bg=COLOR_BLUE_HEX)
    add_body_p(doc, "Buyer acceptance of the 1.5% fee is the single most important thing the pilot must prove. The fallback is a 4% vendor-only commission.", bold_prefix="Key risk: ")

    add_heading_styled(doc, "3. Unit Economics per Released Order", level=1)
    fill_table(doc, ["GHS per released order", "GHS 450 order", "GHS 800 order", "GHS 1,200 order"], [('Croe fee: 2.5% vendor + 1.5% buyer', '+ GHS 18.00', '+ GHS 32.00', '+ GHS 48.00'), ('MoMo collection fee (1.95% of amount collected)', '- GHS 8.91', '- GHS 15.83', '- GHS 23.75'), ('Fraud & loss reserve (0.3% of order)', '- GHS 1.35', '- GHS 2.40', '- GHS 3.60'), ('MoMo payout to vendor (flat)', '- GHS 1.00', '- GHS 1.00', '- GHS 1.00'), ('ID verification (amortised)', '- GHS 1.00', '- GHS 1.00', '- GHS 1.00'), ('Refund leakage (5% of orders refunded)', '- GHS 0.52', '- GHS 0.89', '- GHS 1.30'), ('SMS notifications (~8 per order)', '- GHS 0.35', '- GHS 0.35', '- GHS 0.35'), ('Hosting & AI (variable share)', '- GHS 0.20', '- GHS 0.20', '- GHS 0.20'), ('Net per order at 4%', 'GHS 4.67', 'GHS 10.33', 'GHS 16.80'), ('Net margin (% of order)', '1.04%', '1.29%', '1.40%'), ('Net at 2% vendor-only (old pitch)', '- GHS 4.19', '- GHS 5.42', '- GHS 6.83')], header_bg=COLOR_NAVY_HEX)
    add_body_p(doc, "Base average order GHS 800 (blend of the lead vertical and lower-ticket categories; measured in the pilot). ID checks at GHS 1.00/order assume buyers rely on telco KYC on their MoMo wallet; if every new buyer needs a paid check, net falls to GHS 3.60. No E-Levy (repealed 2 Apr 2025).", bold_prefix="Assumptions: ")

    add_heading_styled(doc, "4. Fixed Operating Costs by Stage (GHS per month, base)", level=1)
    fill_table(doc, ["Monthly fixed cost", "Stage 1 · Legal pilot", "Stage 2 · Compliant launch", "Stage 3 · Scale"], [
        ("People (fully loaded, SSNIT 13% + benefits)", "14,500", "136,000 (to 12 staff)", "290,000 avg (to ~25)"),
        ("Infrastructure & software", "2,500", "20,000", "58,000"),
        ("Office / coworking", "1,500", "15,000", "35,000"),
        ("Vendor acquisition & marketing", "4,000", "40,000", "120,000"),
        ("Legal, compliance, insurance, audit, directors", "-", "21,500", "39,000"),
        ("Travel, admin, misc.", "2,000", "6,000", "15,000"),
        ("Total", "24,500 (US$2,100)", "238,500 (US$20,600)", "557,000 avg (US$48,000)"),
    ], header_bg=COLOR_BLUE_HEX)
    add_body_p(doc, "The BoG licensing pack names five key management roles (CEO, Technology & Systems Manager, Compliance & Risk Manager, Finance Manager, AML Reporting Officer) and requires at least three directors. At month 36, with payroll fully ramped, fixed cost is GHS 637,000/month.", bold_prefix="Why the team is not optional: ")

    add_heading_styled(doc, "5. Capital Required by Stage", level=1)
    fill_table(doc, ["Stage", "Months", "Net need (base)", "Range", "Gate to next stage"], [
        ("1 · Legal pilot", "1-6", "US$26k (GHS 305k)", "US$11k - 51k", "Written aggregator/partner yes; measured buyer drop-off, order value, dispute rate"),
        ("2 · Compliant launch", "7-18", "US$360k (GHS 4.18M)", "US$178k - 662k", "Clean daily reconciliation, disputes < 8%, vendor retention"),
        ("3 · Scale", "19-36", "US$916k (GHS 10.63M)", "US$442k - 1.80M", "Break-even in sight via fee and rails levers"),
        ("Total", "36", "US$1.30M (GHS 15.11M)", "US$631k - 2.52M", ""),
    ], header_bg=COLOR_NAVY_HEX)
    add_body_p(doc, "Pre-seed = stages 1-2: US$387k base (US$189k-713k). Seed of ~US$916k around month 15, on stage 2 data.", bold_prefix="Raise structure: ")

    add_heading_styled(doc, "6. Itemized Costs (GHS, whole stage)", level=1)
    for title, rows in [('Stage 1 - Legal pilot (6 months)', [('Company registration, stamp duty, company secretary', 'one-time', '1,100', '3,000', '5,000'), ('Data Protection Commission registration', 'one-time', '1,000', '2,000', '4,000'), ('Trademark "Croe" (Ghana)', 'one-time', '0', '3,000', '6,000'), ('Apple ($99/yr) + Google Play ($25) developer accounts', 'one-time', '1,440', '1,440', '1,440'), ('Fintech lawyer: custody opinion, ToS, privacy, vendor terms, sandbox/aggregator filings', 'one-time', '35,000', '60,000', '120,000'), ('Pre-pilot security review (automated scan + light external test)', 'one-time', '0', '12,000', '35,000'), ('Test devices (3 Android tiers + 1 iPhone)', 'one-time', '8,000', '15,000', '25,000'), ('Founder stipends (1-2 founders)', 'GHS 10,000/mo x 6', '30,000', '60,000', '120,000'), ('Part-time ops: support, L3 disputes, KYC, daily reconciliation', 'GHS 3,000/mo x 6', '0', '18,000', '30,000'), ('Bookkeeping + company secretary', 'GHS 1,500/mo x 6', '4,800', '9,000', '15,000'), ('Hosting, SMS base, monitoring, tools', 'GHS 2,500/mo x 6', '6,000', '15,000', '24,000'), ('Vendor onboarding and pilot marketing', 'GHS 4,000/mo x 6', '9,000', '24,000', '48,000'), ('Coworking / registered address', 'GHS 1,500/mo x 6', '3,000', '9,000', '18,000'), ('Transport, data, phones, misc.', 'GHS 2,000/mo x 6', '6,000', '12,000', '18,000'), ('Subtotal', '', '105,340', '243,440', '469,440'), ('Contingency 15%', '', '15,801', '36,516', '70,416'), ('Working capital (float + reserve)', '', '10,000', '30,000', '60,000'), ('Less: order income (high case at half volume)', '', '-5,268', '-5,268', '-2,634'), ('Net capital need (GHS)', '', '125,873', '304,688', '597,222'), ('Net capital need (US$)', '', '$10,851', '$26,266', '$51,485')]), ('Stage 2 - Compliant launch (12 months)', [('BoG PFTSP licence: processing GHS 10k + licence GHS 20k', 'one-time', '30,000', '30,000', '30,000'), ('Legal: licence application, partner trust-account agreement, employment, ESOP', 'one-time', '80,000', '130,000', '250,000'), ('Penetration test: web + API + iOS + Android, with retest', 'one-time', '70,000', '140,000', '300,000'), ('ISO 27001 (gap assessment only in Low)', 'one-time', '40,000', '230,000', '520,000'), ('ICT, BCP/DR, AML/CFT and risk policies required by BoG', 'one-time', '20,000', '45,000', '90,000'), ('Custody partner onboarding and integration', 'one-time', '10,000', '40,000', '100,000'), ('Laptops and equipment for new hires', 'one-time', '60,000', '100,000', '150,000'), ('Recruitment (compliance lead, senior engineers)', 'one-time', '20,000', '60,000', '120,000'), ('Payroll, fully loaded, ramping to 12 people', 'GHS 136,000/mo x 12', '1,020,000', '1,632,000', '2,520,000'), ('Infrastructure and software (HA database, GPU box, monitoring, workspace)', 'GHS 20,000/mo x 12', '144,000', '240,000', '384,000'), ('Office / coworking in Accra', 'GHS 15,000/mo x 12', '72,000', '180,000', '300,000'), ('Vendor acquisition and marketing', 'GHS 40,000/mo x 12', '240,000', '480,000', '960,000'), ('Legal and compliance retainer', 'GHS 10,000/mo x 12', '60,000', '120,000', '240,000'), ('Insurance: professional indemnity, cyber, fidelity', 'GHS 5,000/mo x 12', '36,000', '60,000', '120,000'), ('External audit and tax filing', 'GHS 3,500/mo x 12', '24,000', '42,000', '72,000'), ('Directors fees', 'GHS 3,000/mo x 12', '0', '36,000', '96,000'), ('Travel, admin, misc.', 'GHS 6,000/mo x 12', '36,000', '72,000', '120,000'), ('Subtotal', '', '1,962,000', '3,637,000', '6,372,000'), ('Contingency 15%', '', '294,300', '545,550', '955,800'), ('Working capital (float + reserve)', '', '50,000', '233,000', '467,000'), ('Less: order income (high case at half volume)', '', '-235,749', '-235,749', '-117,874'), ('Net capital need (GHS)', '', '2,070,551', '4,179,801', '7,676,926'), ('Net capital need (US$)', '', '$178,496', '$360,328', '$661,804')]), ('Stage 3 - Scale (18 months)', [('ISO 27001 surveillance audits', 'one-time', '60,000', '100,000', '180,000'), ('Annual penetration tests (x2)', 'one-time', '120,000', '250,000', '500,000'), ('Equipment for new hires', 'one-time', '120,000', '170,000', '250,000'), ('Recruitment', 'one-time', '50,000', '120,000', '250,000'), ('Payroll, fully loaded, ramping to about 25 people', 'GHS 290,000/mo x 18', '3,600,000', '5,220,000', '7,560,000'), ('Infrastructure and software (DR region, 2 GPU nodes, log retention, WAF)', 'GHS 58,000/mo x 18', '630,000', '1,044,000', '1,620,000'), ('Office', 'GHS 35,000/mo x 18', '360,000', '630,000', '1,080,000'), ('Vendor acquisition and marketing', 'GHS 120,000/mo x 18', '1,080,000', '2,160,000', '4,500,000'), ('Legal and compliance retainer', 'GHS 15,000/mo x 18', '180,000', '270,000', '540,000'), ('Insurance', 'GHS 10,000/mo x 18', '108,000', '180,000', '360,000'), ('External audit and tax', 'GHS 6,000/mo x 18', '72,000', '108,000', '180,000'), ('Directors fees (incl. independent)', 'GHS 8,000/mo x 18', '54,000', '144,000', '270,000'), ('Travel, admin, misc.', 'GHS 15,000/mo x 18', '144,000', '270,000', '450,000'), ('Subtotal', '', '6,578,000', '10,666,000', '17,740,000'), ('Contingency 15%', '', '986,700', '1,599,900', '2,661,000'), ('Working capital (float + reserve)', '', '300,000', '1,100,000', '1,900,000'), ('Less: order income (high case at half volume)', '', '-2,738,858', '-2,738,858', '-1,369,429'), ('Net capital need (GHS)', '', '5,125,842', '10,627,042', '20,931,571'), ('Net capital need (US$)', '', '$441,883', '$916,124', '$1,804,446')])]:
        add_heading_styled(doc, title, level=2)
        fill_table(doc, ["Item", "Basis (base)", "Low", "Base", "High"], rows)

    add_heading_styled(doc, "7. Working Capital & Float", level=1)
    add_body_p(doc, "Aggregators settle T+1. Escrow holds most orders longer than that, so the float covers same-day deliveries confirmed before settlement. A ring-fenced loss reserve covers wrong adjudications and fraud. Raised and recycled, not burned.")
    fill_table(doc, ["At the end of", "Basis", "Low (GHS)", "Base (GHS)", "High (GHS)"], [
        ("Stage 1 · pilot", "Loss reserve + small float", "10,000", "30,000", "60,000"),
        ("Stage 2 · ~5,000 orders/mo", "One day of order value + partner reserve", "50,000", "233,000", "467,000"),
        ("Stage 3 · ~30,000 orders/mo", "One day of order value + reserve", "300,000", "1,100,000", "1,900,000"),
    ])

    add_heading_styled(doc, "8. Break-Even Sensitivity (orders/month)", level=1)
    add_body_p(doc, "Released orders per month needed to cover the month-36 fixed cost of GHS 637,000/month. All-in fee = 1.5% buyer + the rest from the vendor.")
    fill_table(doc, ["Average order", "3.0% fee", "3.5% fee", "4.0% fee", "4.5% fee", "5.0% fee"], [('GHS 450', 'not viable', 'not viable', '136,300', '92,000', '69,500'), ('GHS 800', 'not viable', '100,600', '61,700', '44,500', '34,800'), ('GHS 1,200', '132,800', '59,000', '37,900', '27,900', '22,100'), ('GHS 2,000', '65,500', '32,300', '21,400', '16,000', '12,800')], header_bg=COLOR_NAVY_HEX)

    add_heading_styled(doc, "9. Levers and Risks", level=1)
    fill_table(doc, ["", "Change", "Effect (GHS 800 order, month-36 cost base)"], [
        ("Lever · pricing", "5% fee (3.5% vendor + 1.5% buyer)", "Net GHS 18.33 -> ~34,800 orders/month"),
        ("Lever · own the rails", "Direct telco collection at ~1% (needs PSP licence: Medium GHS 0.8M / Enhanced GHS 2M capital)", "Net GHS 18.45 -> ~34,500 orders/month"),
        ("Lever · ticket size", "GHS 1,200 average order", "Net GHS 16.80 -> ~37,900 orders/month"),
        ("Risk · buyer ID checks", "Every new buyer needs a paid ID check", "Net falls to GHS 3.60 per order"),
        ("Risk · partner fee", "Custody bank takes 0.25% of volume", "Net GHS 8.33 -> ~76,500 orders/month"),
        ("Risk · regulatory delay", "Licence or sandbox slips 3-6 months", "GHS 0.72M - 1.43M extra stage 2 burn"),
    ], header_bg=COLOR_BLUE_HEX)
    add_callout(doc, "At month 36 the plan's contribution is GHS 310k/month against GHS 637k of fixed cost, a gap of ~GHS 327k/month (US$28k). Closing it takes the pricing or rails lever, or another 6-12 months of runway.", title="THE MONTH-36 GAP, STATED PLAINLY")

    add_heading_styled(doc, "10. Pilot Success Thresholds (Stage 1 Gate)", level=1)
    add_body_p(doc, "1. Buyer fee acceptance: checkout drop-off below 30% with the 1.5% fee shown.\n2. At least 250 released orders in month 6.\n3. At least GHS 10.00 net per order at the measured order value.\n4. Dispute rate below 8%.\n5. Over 95% of claims resolved without a successful appeal.\n6. GHS 0.00 variance across daily three-way reconciliation.")

    add_heading_styled(doc, "11. Sources & Confidence", level=1)
    fill_table(doc, ["Input", "Value used", "Status"], [
        ("FX", "US$1 = GHS 11.60 (21-25 Sep 2026)", "Published"),
        ("MoMo collection / payout", "1.95% / GHS 1.00 per transfer (Paystack; Hubtel)", "Published"),
        ("BoG licence fees & capital", "PFTSP: no capital, GHS 10k + 20k; Medium GHS 0.8M; Enhanced GHS 2M; DEMI GHS 20M", "Confirm (2020 table)"),
        ("Key management roles", "CEO, Tech & Systems, Compliance & Risk, Finance, AMLRO; 3+ directors", "Published"),
        ("Company registration", "GHS 585 + 1% stamp duty (ORC, Feb 2026)", "Published"),
        ("Salaries / on-costs", "Senior SWE Accra GHS 139k-504k/yr (levels.fyi); SSNIT 13%", "Estimate / Published"),
        ("GPU / hosting", "Hetzner GEX44 EUR 184-234/mo; Render web $25, Postgres from $6", "Published"),
        ("Pentest / ISO 27001", "US$6k-26k / US$15k-45k first year", "Published guides"),
        ("ID verification", "US$0.50-2.00 per onboarding (Smile ID, sales-quoted)", "Confirm"),
        ("Legal, insurance, office, directors, partner fees", "See section 6", "Estimate"),
        ("Order value, refund rate, loss reserve, volume", "GHS 800; 5%; 0.3%; 5k/mo at month 18, 30k at month 36", "Confirm in pilot"),
    ])
    add_body_p(doc, "Full source links: pitch/FINANCIAL-MODEL.md section 12. Not legal, tax or investment advice.", space_after=4)

    add_footer(doc, "Financial Model & Unit Economics")
    out_file = os.path.join(PITCH_DIR, "Croe_Financial_Model.docx")
    doc.save(out_file)
    print(f"Generated: {out_file}")

# ==============================================================================
# DOC 3: Banking Partner & Custody Brief (.docx)
# ==============================================================================
def generate_banking_brief_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

    add_header_banner(
        doc,
        title="croe. - Banking Partner & Custody Brief",
        subtitle="Institutional Due-Diligence, Trust Architecture & Regulatory Compliance",
        doc_ref="INSTITUTIONAL PARTNERSHIP BRIEF · BOG ACT 987 COMPLIANT"
    )

    add_heading_styled(doc, "1. Executive Summary & Partnership Proposition", level=1)
    add_body_p(doc, "Croe is a mobile-first trust and escrow transaction service for emerging-market social commerce, launching Ghana-first on Mobile Money rails. Croe acts as a technology orchestration platform that channels high-velocity social commerce payments into a licensed partner bank trust account.", bold_prefix="The Opportunity: ")

    t_prop = doc.add_table(rows=3, cols=2)
    for i, h in enumerate(["Value Delivered to Partner Bank", "Operational Role of Partner Bank"]):
        t_prop.cell(0, i).paragraphs[0].text = h
    rows_prop = [
        ("Low-Cost Float Growth: Sticky, non-interest-bearing escrow deposits residing in a pooled trust account (CASA float).", "Custodial Trust Account: Provide a dedicated pooled escrow trust account under Bank of Ghana guidelines."),
        ("Transactional Revenue: Transaction fee revenue on high-velocity inbound collections and outbound B2C wallet disbursements.", "Automated Reconciliation: Automated electronic statement feed (MT940 / API) for daily 3-way balance matching.")
    ]
    for r_idx, r_data in enumerate(rows_prop, start=1):
        for c_idx, val in enumerate(r_data):
            t_prop.cell(r_idx, c_idx).paragraphs[0].text = val
    style_table(t_prop, header_bg=COLOR_NAVY_HEX)

    add_heading_styled(doc, "2. Regulatory Architecture: The Non-Custodial Tech Provider", level=1)
    add_body_p(doc, "Croe is not a bank, does not issue electronic money, and never holds customer funds on its corporate balance sheet. Holding third-party funds in Ghana is governed by the Bank of Ghana under the Payment Systems and Services Act, 2019 (Act 987). Croe achieves compliance via a Phased Custody Architecture where float is held strictly by licensed institutions:", bold_prefix="Act 987 Compliance: ")

    t_reg = doc.add_table(rows=4, cols=4)
    for i, h in enumerate(["Phase", "Legal Role of Croe", "Float Holder", "Compliance Standing"]):
        t_reg.cell(0, i).paragraphs[0].text = h
    rows_reg = [
        ("P0: Build", "Software Developer", "Nobody (Sandbox)", "Development only; 354 automated tests."),
        ("P1: Pilot", "Technology Merchant", "Aggregator Settlement Account", "Interim supervised pilot; capped volumes (~250 orders/month by month 6)."),
        ("P2: Partner-Held", "Payment Tech Provider (PFTSP)", "Licensed Bank Pooled Trust Account", "Fully compliant institutional launch; bank-held trust account.")
    ]
    for r_idx, r_data in enumerate(rows_reg, start=1):
        for c_idx, val in enumerate(r_data):
            t_reg.cell(r_idx, c_idx).paragraphs[0].text = val
    style_table(t_reg, header_bg=COLOR_BLUE_HEX)

    add_heading_styled(doc, "3. Double-Spend Defense & Ledger Immutability", level=1)
    add_body_p(doc, "Croe incorporates three structural layers of defense preventing financial double-spend or unauthorized disbursement:", bold_prefix="Triple-Lock Concurrency: ")
    add_body_p(doc, "1. Redis Distributed Fast Dedup: SETNX with 24-hour TTL guarantees instantaneous short-circuiting of duplicate webhooks.\n2. Row Locking (SELECT ... FOR UPDATE): All state mutations serialize concurrent requests at the PostgreSQL database row level.\n3. Partial Unique Indexes: Error 23505 traps guarantee that concurrent deposits or duplicate payouts fail structurally.")
    add_body_p(doc, "Balance history resides in an append-only transaction_ledger where UPDATE and DELETE privileges are permanently revoked at the database role level.", bold_prefix="Append-Only Sub-Ledger: ")

    add_heading_styled(doc, "4. KYC & AML Compliance Controls", level=1)
    t_kyc = doc.add_table(rows=4, cols=4)
    for i, h in enumerate(["Tier", "Verification Requirements", "Per-Transaction Cap", "Daily Limit"]):
        t_kyc.cell(0, i).paragraphs[0].text = h
    rows_kyc = [
        ("Tier 0", "Phone verified via MoMo OTP", "GHS 500.00", "GHS 1,000.00"),
        ("Tier 1", "Full Name + Ghana Card / Passport", "GHS 5,000.00", "GHS 20,000.00"),
        ("Tier 2", "Enhanced Due Diligence (EDD) + Address", "GHS 50,000.00", "GHS 100,000.00")
    ]
    for r_idx, r_data in enumerate(rows_kyc, start=1):
        for c_idx, val in enumerate(r_data):
            t_kyc.cell(r_idx, c_idx).paragraphs[0].text = val
    style_table(t_kyc)

    add_heading_styled(doc, "5. Automated Daily Three-Way Reconciliation", level=1)
    add_body_p(doc, "Every 24 hours, Croe executes an automated three-way reconciliation matching: (1) Bank Partner Escrow Float Account statement balance, (2) Sum of active unreleased funds in the immutable transaction ledger, and (3) Physical payment rail transaction logs. If any non-zero variance occurs, the system logs a Priority 1 (P1) incident, halts automated disbursement workers, and alerts operations.", bold_prefix="Circuit Breaker: ")

    add_heading_styled(doc, "6. Projected Commercial Scale for Banking Partner", level=1)
    t_proj = doc.add_table(rows=4, cols=4)
    for i, h in enumerate(["Milestone", "Monthly Completed Txns", "Monthly Processing Volume", "Average Daily Float in Bank"]):
        t_proj.cell(0, i).paragraphs[0].text = h
    rows_proj = [
        ("Supervised Pilot (month 6)", "250 txns / mo", "GHS 200,000.00", "GHS 13,000 – 20,000"),
        ("Compliant Launch (month 18)", "5,000 txns / mo", "GHS 4,000,000.00", "GHS 267,000 – 400,000"),
        ("Scale (month 36)", "30,000 txns / mo", "GHS 24,000,000.00", "GHS 1.6M – 2.4M")
    ]
    for r_idx, r_data in enumerate(rows_proj, start=1):
        for c_idx, val in enumerate(r_data):
            t_proj.cell(r_idx, c_idx).paragraphs[0].text = val
    style_table(t_proj)

    add_footer(doc, "Banking Partner & Custody Brief")
    out_file = os.path.join(PITCH_DIR, "Croe_Banking_Partner_Brief.docx")
    doc.save(out_file)
    print(f"Generated: {out_file}")

# ==============================================================================
# DOC 4: Vendor Pilot Guide & Agreement (.docx)
# ==============================================================================
def generate_vendor_agreement_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

    add_header_banner(
        doc,
        title="croe. - Vendor Pilot Guide & Agreement",
        subtitle="How Croe Closes Your Sales & Pilot Onboarding Agreement",
        doc_ref="VENDOR PILOT COHORT · ACCRA, GHANA"
    )

    add_heading_styled(doc, "Part 1: How Croe Works for Your Shop", level=1)
    add_body_p(doc, "Stop losing qualified customers to the 'Send the MoMo first' standoff. When selling smartphones, sneakers, or fashion on WhatsApp and Instagram, buyers refuse to pay upfront due to scam fears, and you cannot dispatch without payment. Croe solves this standoff by holding the buyer's money securely until delivery is confirmed, then paying your MoMo wallet automatically.", bold_prefix="The Conversion Pitch: ")

    t_steps = doc.add_table(rows=5, cols=3)
    for i, h in enumerate(["Step", "What You Do", "What the Buyer Experiences"]):
        t_steps.cell(0, i).paragraphs[0].text = h
    rows_steps = [
        ("1. Create Link", "Generate escrow link in Croe app (10 sec)", "Receives link in WhatsApp/IG DM with item details."),
        ("2. Buyer Pays", "You see funds are locked and guaranteed", "Authorizes payment via standard MoMo prompt (MTN, Telecel, AirtelTigo)."),
        ("3. Dispatch", "Send your dispatch rider with peace of mind", "Tracks delivery knowing their money is protected."),
        ("4. Cashout", "Delivery confirmed -> Instant MoMo payout", "Taps 'Received' in browser/app -> MoMo hits your wallet.")
    ]
    for r_idx, r_data in enumerate(rows_steps, start=1):
        for c_idx, val in enumerate(r_data):
            t_steps.cell(r_idx, c_idx).paragraphs[0].text = val
    style_table(t_steps, header_bg=COLOR_NAVY_HEX)

    add_heading_styled(doc, "Transparent Pricing: 2.5% per Completed Sale", level=1)
    t_fee = doc.add_table(rows=6, cols=3)
    for i, h in enumerate(["Feature", "Croe Escrow", "Traditional Way (Direct MoMo)"]):
        t_fee.cell(0, i).paragraphs[0].text = h
    rows_fee = [
        ("Setup Fee", "GHS 0.00 (Free)", "GHS 0.00"),
        ("Monthly Subscription", "GHS 0.00 (Free)", "GHS 0.00"),
        ("Fee on Completed Sale", "2.5% (e.g. GHS 30 on a GHS 1,200 order)", "0% (but you lose 25% of sales)"),
        ("Buyer Protection Fee", "1.5%, paid by the buyer at checkout", "-"),
        ("Cancelled / Refunded Order", "GHS 0.00 (Free)", "Wasted rider fees, lost stock")
    ]
    for r_idx, r_data in enumerate(rows_fee, start=1):
        for c_idx, val in enumerate(r_data):
            t_fee.cell(r_idx, c_idx).paragraphs[0].text = val
    style_table(t_fee, header_bg=COLOR_BLUE_HEX)

    add_heading_styled(doc, "Part 2: Non-Binding Letter of Intent (LOI)", level=1)
    add_body_p(doc, "This is a non-binding statement of intent for participation in the Croe 60-day vendor pilot. It does not create financial or legal obligations on either party.")

    add_body_p(doc, "If Croe launches and functions as described, I intend to use it for my online sales and accept the standard commission of 2.5% on completed, successfully released orders. I understand my buyers would pay a separate 1.5% protection fee at checkout.", bold_prefix="Statement of Intent: ")

    t_form = doc.add_table(rows=8, cols=2)
    for i, h in enumerate(["Vendor Detail", "Response"]):
        t_form.cell(0, i).paragraphs[0].text = h
    rows_form = [
        ("Business / Trading Name", "__________________________________________________"),
        ("Goods Sold (Phones / Sneakers / Fashion)", "__________________________________________________"),
        ("Instagram / TikTok Handle", "__________________________________________________"),
        ("WhatsApp Contact Number", "__________________________________________________"),
        ("Approximate Orders per Month", "__________________________________________________"),
        ("Average Order Value (GHS)", "__________________________________________________"),
        ("Would your buyers accept a 1.5% protection fee? (Yes / No / Not sure)", "__________________________________________________")
    ]
    for r_idx, r_data in enumerate(rows_form, start=1):
        for c_idx, val in enumerate(r_data):
            t_form.cell(r_idx, c_idx).paragraphs[0].text = val
    style_table(t_form)

    add_body_p(doc, "Signature: ___________________________          Date: ___________________________", space_after=12)
    add_body_p(doc, "Representative Name: ___________________________", space_after=16)

    add_footer(doc, "Vendor Pilot Guide & Agreement")
    out_file = os.path.join(PITCH_DIR, "Croe_Vendor_Pilot_Agreement.docx")
    doc.save(out_file)
    print(f"Generated: {out_file}")

if __name__ == "__main__":
    generate_one_pager_docx()
    generate_financial_model_docx()
    generate_banking_brief_docx()
    generate_vendor_agreement_docx()
    print("All Word documents generated successfully!")
